import csv
import io
from datetime import date
from decimal import Decimal

from django.db import transaction
from django.http import HttpResponse
from openpyxl import Workbook, load_workbook

from apps.academics.models import AcademicPeriod, AcademicPolicy, AcademicYear, Classroom
from .models import Enrollment, Guardian, Student, Teacher
from .serializers import GuardianSerializer, StudentSerializer, TeacherSerializer


HEADER_ALIASES = {
    "first_name": {
        "first_name", "firstname", "prenom", "prénom", "given_name",
    },
    "last_name": {
        "last_name", "lastname", "nom", "surname", "family_name",
    },
    "matricule": {
        "matricule", "student_id", "student_number", "numero", "numéro",
    },
    "employee_number": {
        "employee_number", "teacher_id", "teacher_number",
        "matricule_enseignant", "matricule enseignant",
    },
    "gender": {"gender", "sexe"},
    "date_of_birth": {
        "date_of_birth", "birth_date", "date_naissance", "date de naissance",
    },
    "place_of_birth": {
        "place_of_birth", "birth_place", "lieu_naissance", "lieu de naissance",
    },
    "nationality": {"nationality", "nationalite", "nationalité"},
    "address": {"address", "adresse"},
    "phone": {"phone", "telephone", "téléphone", "tel"},
    "alternate_phone": {
        "alternate_phone", "secondary_phone", "telephone_2", "téléphone 2",
    },
    "email": {"email", "e-mail", "mail"},
    "admission_date": {
        "admission_date", "date_admission", "date d'admission",
    },
    "speciality": {
        "speciality", "specialty", "specialite", "spécialité",
    },
    "hire_date": {
        "hire_date", "date_embauche", "date d'embauche",
    },
    "occupation": {"occupation", "profession"},
    "preferred_language": {
        "preferred_language", "language", "langue",
    },
    "status": {"status", "statut"},
    "notes": {"notes", "note", "comment", "commentaire"},
}


STUDENT_FIELDS = [
    "matricule",
    "first_name",
    "last_name",
    "gender",
    "date_of_birth",
    "place_of_birth",
    "nationality",
    "address",
    "phone",
    "email",
    "admission_date",
    "status",
    "notes",
]

TEACHER_FIELDS = [
    "employee_number",
    "first_name",
    "last_name",
    "phone",
    "email",
    "speciality",
    "hire_date",
    "status",
    "notes",
]

GUARDIAN_FIELDS = [
    "first_name",
    "last_name",
    "phone",
    "alternate_phone",
    "email",
    "occupation",
    "address",
    "preferred_language",
    "is_active",
]


def normalize_header(value):
    value = str(value or "").strip().lower().replace("-", "_")
    for canonical, aliases in HEADER_ALIASES.items():
        if value in aliases:
            return canonical
    return value.replace(" ", "_")


def normalize_value(value):
    if value is None:
        return ""
    if isinstance(value, (date,)):
        return value.isoformat()
    return str(value).strip()


def read_csv(upload):
    raw = upload.read()
    text = raw.decode("utf-8-sig")

    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;|\t")
    except csv.Error:
        dialect = csv.excel

    reader = csv.DictReader(io.StringIO(text), dialect=dialect)

    rows = []
    for raw_row in reader:
        rows.append({
            normalize_header(key): normalize_value(value)
            for key, value in raw_row.items()
            if key is not None
        })
    return rows


def read_xlsx(upload):
    workbook = load_workbook(upload, read_only=True, data_only=True)
    sheet = workbook.active

    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator, None)

    if not headers:
        return []

    headers = [normalize_header(value) for value in headers]

    rows = []
    for values in iterator:
        row = {
            headers[index]: normalize_value(value)
            for index, value in enumerate(values)
            if index < len(headers) and headers[index]
        }
        if any(value for value in row.values()):
            rows.append(row)

    return rows


def read_upload(upload):
    filename = (upload.name or "").lower()

    if filename.endswith(".csv"):
        return read_csv(upload)

    if filename.endswith(".xlsx"):
        return read_xlsx(upload)

    raise ValueError("Format non supporté. Utilisez CSV ou XLSX.")


def parse_boolean(value, default=True):
    if value in ("", None):
        return default
    return str(value).strip().lower() in {"1", "true", "yes", "oui", "y"}


def normalize_student(row):
    payload = {key: row.get(key, "") for key in STUDENT_FIELDS}

    if payload.get("gender"):
        gender = payload["gender"].strip().upper()
        aliases = {
            "M": "MALE",
            "MALE": "MALE",
            "MASCULIN": "MALE",
            "HOMME": "MALE",
            "F": "FEMALE",
            "FEMALE": "FEMALE",
            "FEMININ": "FEMALE",
            "FÉMININ": "FEMALE",
            "FEMME": "FEMALE",
        }
        payload["gender"] = aliases.get(gender, gender)

    if payload.get("status"):
        payload["status"] = payload["status"].upper()
    else:
        payload["status"] = Student.Status.ACTIVE

    payload["date_of_birth"] = payload["date_of_birth"] or None
    payload["admission_date"] = payload["admission_date"] or None

    return payload


def normalize_teacher(row):
    payload = {key: row.get(key, "") for key in TEACHER_FIELDS}
    payload["status"] = (
        payload["status"].upper()
        if payload.get("status")
        else Teacher.Status.ACTIVE
    )
    payload["hire_date"] = payload["hire_date"] or None
    return payload


def normalize_guardian(row):
    payload = {key: row.get(key, "") for key in GUARDIAN_FIELDS}
    payload["preferred_language"] = (
        payload.get("preferred_language", "FR").strip().upper() or "FR"
    )
    if payload["preferred_language"] in {"FRENCH", "FRANCAIS", "FRANÇAIS"}:
        payload["preferred_language"] = "FR"
    if payload["preferred_language"] in {"ENGLISH", "ANGLAIS"}:
        payload["preferred_language"] = "EN"

    payload["is_active"] = parse_boolean(
        row.get("is_active", ""),
        default=True,
    )
    return payload


def serializer_for_entity(entity):
    if entity == "students":
        return StudentSerializer, normalize_student, Student
    if entity == "teachers":
        return TeacherSerializer, normalize_teacher, Teacher
    if entity == "guardians":
        return GuardianSerializer, normalize_guardian, Guardian
    raise ValueError("Type d'import inconnu.")


def find_existing(entity, school, payload):
    if entity == "students" and payload.get("matricule"):
        return Student.objects.filter(
            school=school,
            matricule__iexact=payload["matricule"],
        ).first()

    if entity == "teachers" and payload.get("employee_number"):
        return Teacher.objects.filter(
            school=school,
            employee_number__iexact=payload["employee_number"],
        ).first()

    if entity == "guardians" and payload.get("phone"):
        return Guardian.objects.filter(
            school=school,
            phone=payload["phone"],
        ).first()

    return None


def import_people(
    *,
    request,
    entity,
    upload,
    dry_run=False,
    academic_year_id=None,
    classroom_id=None,
):
    rows = read_upload(upload)
    serializer_class, normalize, _model = serializer_for_entity(entity)

    school = request.school

    classroom = None
    academic_year = None

    if classroom_id:
        classroom = Classroom.objects.filter(
            school=school,
            id=classroom_id,
        ).select_related("academic_year").first()

        if not classroom:
            raise ValueError("Classe introuvable dans cet établissement.")

        academic_year = classroom.academic_year

    if academic_year_id:
        academic_year = AcademicYear.objects.filter(
            school=school,
            id=academic_year_id,
        ).first()
        if not academic_year:
            raise ValueError("Année scolaire introuvable.")

    if classroom and academic_year and classroom.academic_year_id != academic_year.id:
        raise ValueError("La classe n'appartient pas à l'année scolaire sélectionnée.")

    result = {
        "entity": entity,
        "dry_run": dry_run,
        "total_rows": len(rows),
        "created": 0,
        "updated": 0,
        "enrolled": 0,
        "errors": [],
    }

    for index, row in enumerate(rows, start=2):
        try:
            payload = normalize(row)

            if not payload.get("first_name") or not payload.get("last_name"):
                raise ValueError("Prénom et nom sont obligatoires.")

            existing = find_existing(entity, school, payload)

            serializer = serializer_class(
                existing,
                data=payload,
                partial=bool(existing),
                context={"request": request},
            )
            serializer.is_valid(raise_exception=True)

            if dry_run:
                if existing:
                    result["updated"] += 1
                else:
                    result["created"] += 1
                if entity == "students" and classroom:
                    result["enrolled"] += 1
                continue

            with transaction.atomic():
                instance = serializer.save()

                if existing:
                    result["updated"] += 1
                else:
                    result["created"] += 1

                if entity == "students" and classroom:
                    Enrollment.objects.update_or_create(
                        school=school,
                        student=instance,
                        academic_year=academic_year,
                        defaults={
                            "classroom": classroom,
                            "status": Enrollment.Status.ACTIVE,
                        },
                    )
                    result["enrolled"] += 1

        except Exception as exc:
            detail = getattr(exc, "detail", None)
            if detail is not None:
                message = str(detail)
            else:
                message = str(exc)

            result["errors"].append({
                "row": index,
                "message": message,
            })

    return result


def effective_threshold(level):
    policy, _ = AcademicPolicy.objects.get_or_create(
        school=level.school,
    )

    if level.promotion_threshold_override is not None:
        return level.promotion_threshold_override

    if level.cycle.promotion_threshold_override is not None:
        return level.cycle.promotion_threshold_override

    return policy.default_promotion_threshold


def build_export_rows(entity, school, academic_year_id=None, classroom_id=None):
    if entity == "students":
        queryset = Student.objects.filter(school=school)

        if academic_year_id:
            queryset = queryset.filter(
                enrollments__academic_year_id=academic_year_id
            )

        if classroom_id:
            queryset = queryset.filter(
                enrollments__classroom_id=classroom_id
            )

        queryset = queryset.distinct().order_by("last_name", "first_name")

        headers = [
            "matricule", "last_name", "first_name", "gender",
            "date_of_birth", "phone", "email", "status",
        ]
        rows = [
            [
                item.matricule,
                item.last_name,
                item.first_name,
                item.gender,
                item.date_of_birth or "",
                item.phone,
                item.email,
                item.status,
            ]
            for item in queryset
        ]
        return headers, rows

    if entity == "teachers":
        queryset = Teacher.objects.filter(
            school=school
        ).order_by("last_name", "first_name")

        headers = [
            "employee_number", "last_name", "first_name",
            "phone", "email", "speciality", "status",
        ]
        rows = [
            [
                item.employee_number,
                item.last_name,
                item.first_name,
                item.phone,
                item.email,
                item.speciality,
                item.status,
            ]
            for item in queryset
        ]
        return headers, rows

    if entity == "guardians":
        queryset = Guardian.objects.filter(
            school=school
        ).order_by("last_name", "first_name")

        headers = [
            "last_name", "first_name", "phone", "alternate_phone",
            "email", "occupation", "preferred_language", "is_active",
        ]
        rows = [
            [
                item.last_name,
                item.first_name,
                item.phone,
                item.alternate_phone,
                item.email,
                item.occupation,
                item.preferred_language,
                item.is_active,
            ]
            for item in queryset
        ]
        return headers, rows

    if entity == "enrollments":
        queryset = Enrollment.objects.filter(
            school=school
        ).select_related(
            "student",
            "academic_year",
            "classroom",
            "classroom__level",
        )

        if academic_year_id:
            queryset = queryset.filter(
                academic_year_id=academic_year_id
            )

        if classroom_id:
            queryset = queryset.filter(classroom_id=classroom_id)

        headers = [
            "student_matricule", "student_name", "academic_year",
            "classroom", "level", "status", "final_average",
            "promotion_decision",
        ]
        rows = [
            [
                item.student.matricule,
                f"{item.student.last_name} {item.student.first_name}".strip(),
                item.academic_year.name,
                item.classroom.name,
                item.classroom.level.name,
                item.status,
                item.final_average if item.final_average is not None else "",
                item.promotion_decision,
            ]
            for item in queryset
        ]
        return headers, rows

    raise ValueError("Type d'export inconnu.")


def csv_response(filename, headers, rows):
    response = HttpResponse(
        content_type="text/csv; charset=utf-8",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="{filename}.csv"'
    )
    response.write("\ufeff")

    writer = csv.writer(response)
    writer.writerow(headers)
    writer.writerows(rows)

    return response


def xlsx_response(filename, headers, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"

    sheet.append(headers)

    for row in rows:
        sheet.append(list(row))

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = (
        f'attachment; filename="{filename}.xlsx"'
    )

    return response


def template_response(entity):
    if entity == "students":
        headers = STUDENT_FIELDS
        example = [
            "STU-00001",
            "Kevin",
            "Omgba",
            "MALE",
            "2011-03-12",
            "Yaoundé",
            "Camerounaise",
            "Yaoundé",
            "699000000",
            "student@example.com",
            "2026-09-01",
            "ACTIVE",
            "",
        ]
    elif entity == "teachers":
        headers = TEACHER_FIELDS
        example = [
            "TCH-00001",
            "Alice",
            "Nana",
            "677000000",
            "teacher@example.com",
            "Mathématiques",
            "2024-09-01",
            "ACTIVE",
            "",
        ]
    elif entity == "guardians":
        headers = GUARDIAN_FIELDS
        example = [
            "Marie",
            "Omgba",
            "690000000",
            "",
            "parent@example.com",
            "Commerçante",
            "Yaoundé",
            "FR",
            "TRUE",
        ]
    else:
        raise ValueError("Type de modèle inconnu.")

    return csv_response(
        f"template-{entity}",
        headers,
        [example],
    )
